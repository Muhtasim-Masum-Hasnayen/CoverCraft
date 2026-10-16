from __future__ import annotations

import io

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from services.fonts import DOCX_FONT
from services.pdf_generator import render_pdf
from templates import build_layout

ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}


def _rgb(hexcolor):
    h = hexcolor.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _set_font(run, name, size, bold=False, italic=False, color="#111827"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = _rgb(color)


def _fmt(p, align="left", before=0.0, after=0.0, line=1.12):
    pf = p.paragraph_format
    pf.alignment = ALIGN[align]
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = line
    return p


def _border_xml(tag, val, sz, color, space=0):
    el = OxmlElement(tag)
    el.set(qn("w:val"), val); el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space)); el.set(qn("w:color"), color.lstrip("#"))
    return el


def _table_borders(table, color, sz=4, outer=False, inside_h=True):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in ("top", "left", "right", "insideV") or (edge == "insideH" and not inside_h):
            borders.append(_border_xml(f"w:{edge}", "nil", 0, "FFFFFF"))
        else:
            borders.append(_border_xml(f"w:{edge}", "single", sz, color))
    tblPr.append(borders)


def _cell_bottom_border(cell, color, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    b.append(_border_xml("w:bottom", "single", sz, color))
    tcPr.append(b)


def _fixed_table(doc, rows, widths_cm):
    t = doc.add_table(rows=rows, cols=len(widths_cm))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    tblPr = t._tbl.tblPr
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblPr.append(layout)
    for col, w in zip(t.columns, widths_cm):
        col.width = Cm(w)  # updates <w:tblGrid> so Word and LibreOffice agree
    for row in t.rows:
        for c, w in zip(row.cells, widths_cm):
            c.width = Cm(w)
    return t


def _cell_text(cell, text, font, size, *, bold=False, color="#111827", first=True, after=1):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    _fmt(p, "left", 0, after)
    _set_font(p.add_run(text), font, size, bold=bold, color=color)
    return p


def _spacer(doc, h):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(max(h, 1))
    return p


def generate_docx(data, scale=None):
    spec, blocks = build_layout(data)
    if scale is None:  # reuse the PDF fit-to-one-page result; fonts are metric-compatible
        scale = render_pdf(data)[1]
    font = DOCX_FONT[spec["font"]]
    accent = spec["accent"]
    inst = data["institution"]

    doc = Document()
    doc.core_properties.title = spec["title"]
    doc.core_properties.author = ""
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    margin = 2.54 if spec["border"] != "none" else 2.0
    sec.left_margin = sec.right_margin = Cm(margin)
    sec.top_margin = Cm(margin)
    sec.bottom_margin = Cm(margin + (0.3 if spec["footer"] else 0))
    text_w = 21.0 - 2 * margin

    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(11)

    if spec["border"] != "none":
        pg = OxmlElement("w:pgBorders"); pg.set(qn("w:offsetFrom"), "page")
        val, sz = ("double", 12) if spec["border"] == "double" else ("single", 8)
        for edge in ("top", "left", "bottom", "right"):
            pg.append(_border_xml(f"w:{edge}", val, sz, accent, space=26))
        sec._sectPr.append(pg)

    if spec["footer"]:
        fp = sec.footer.paragraphs[0]
        _fmt(fp, "center")
        _set_font(fp.add_run(spec["footer"]), font, 9, color="#6B7280")

    for b in blocks:
        t = b["t"]
        if t == "spacer":
            _spacer(doc, b["h"] * scale)
        elif t == "text":
            text = b["text"].upper() if b["caps"] else b["text"]
            p = _fmt(doc.add_paragraph(), b["align"], b["before"] * scale, b["after"] * scale)
            _set_font(p.add_run(text), font, b["size"] * scale, b["bold"], b["italic"], b["color"])
        elif t == "rule":
            p = doc.add_paragraph()
            _fmt(p, b["align"], b["before"] * scale, b["after"] * scale, line=1.0)
            r = p.add_run(""); _set_font(r, font, 2)
            pPr = p._p.get_or_add_pPr()
            bd = OxmlElement("w:pBdr")
            bd.append(_border_xml("w:bottom", "single", max(2, int(b["thickness"] * 8 * scale)), b["color"]))
            pPr.append(bd)
            gap = text_w * (1 - b["width"])
            if b["align"] == "center":
                p.paragraph_format.left_indent = Cm(gap / 2); p.paragraph_format.right_indent = Cm(gap / 2)
            else:
                p.paragraph_format.right_indent = Cm(gap)
        elif t == "logo" and inst["logo"]:
            w_px, h_px = inst["logo_size"]
            h = b["max_h"] * scale
            w = h * w_px / h_px
            if w > text_w * 28.35 * 0.5:
                w = text_w * 28.35 * 0.5; h = w * h_px / w_px
            p = _fmt(doc.add_paragraph(), b["align"], 0, b["after"] * scale, line=1.0)
            p.add_run().add_picture(io.BytesIO(inst["logo"]), width=Pt(w), height=Pt(h))
        elif t == "columns":
            size = b["size"] * scale
            gap = 0.8
            cw = (text_w - gap) / 2
            tbl = _fixed_table(doc, 1, [cw, gap, cw])
            _table_borders(tbl, "FFFFFF", inside_h=False)
            for ci, col in zip((0, 2), b["cols"]):
                cell = tbl.rows[0].cells[ci]
                _cell_text(cell, col["heading"].upper(), font, size - 1.5 * scale, bold=True, color=accent, after=3)
                bp = cell.paragraphs[0]._p.get_or_add_pPr()
                bd = OxmlElement("w:pBdr"); bd.append(_border_xml("w:bottom", "single", 6, accent, space=2)); bp.append(bd)
                for text, bold in col["lines"]:
                    _cell_text(cell, text, font, size, bold=bold, first=False)
        elif t == "members":
            size = b["size"] * scale
            ms = b["members"]
            has_role = any(m["role"] for m in ms)
            head = ["#", "Name", "Student ID"] + (["Role"] if has_role else [])
            fr = [0.07, 0.43, 0.25, 0.25] if has_role else [0.08, 0.55, 0.37]
            tbl = _fixed_table(doc, len(ms) + 1, [text_w * f for f in fr])
            _table_borders(tbl, "D1D5DB", sz=4)
            for ci, h in enumerate(head):
                cell = tbl.rows[0].cells[ci]
                _cell_text(cell, h, font, size - 1 * scale, bold=True, color=accent)
                _cell_bottom_border(cell, accent, 8)
            for ri, m in enumerate(ms, 1):
                vals = [str(ri), m["name"], m["student_id"]] + ([m["role"]] if has_role else [])
                for ci, v in enumerate(vals):
                    _cell_text(tbl.rows[ri].cells[ci], v, font, size)
        elif t == "kv":
            size = b["size"] * scale
            tbl = _fixed_table(doc, len(b["rows"]), [text_w * 0.28, text_w * 0.72])
            _table_borders(tbl, "E5E7EB", sz=4)
            for ri, (k, v) in enumerate(b["rows"]):
                _cell_text(tbl.rows[ri].cells[0], k.upper(), font, size - 2 * scale, bold=True, color=accent, after=2)
                _cell_text(tbl.rows[ri].cells[1], v, font, size, after=2)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
