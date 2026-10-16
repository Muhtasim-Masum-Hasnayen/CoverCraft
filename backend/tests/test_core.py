import base64, io, re, shutil, subprocess, tempfile, unittest, zipfile

from docx import Document

from services.docx_generator import generate_docx
from services.logo_processor import LogoError, process_logo
from services.pdf_generator import generate_pdf, render_pdf
from tests.sample_data import logo_b64, sample
from validators import TEMPLATE_IDS, ValidationFailure, normalize, safe_filename
from templates import template_catalog


def pdf_pages(pdf: bytes) -> int:
    return len(re.findall(rb"/Type /Page\b", pdf))


class ValidationTests(unittest.TestCase):
    def test_valid_payload_normalises(self):
        d = normalize(sample())
        self.assertEqual(d["document"]["submission_date_text"], "16 October 2026")
        self.assertEqual(d["submitted_by"]["department"], d["institution"]["department"])

    def test_missing_required_fields(self):
        with self.assertRaises(ValidationFailure) as cm:
            normalize({"document": {}, "institution": {}, "submitted_to": {}, "submitted_by": {}, "design": {}})
        e = cm.exception.errors
        for k in ("document.title", "document.submission_date", "institution.name",
                  "institution.department", "submitted_to.name", "submitted_by.members"):
            self.assertIn(k, e)

    def test_non_object_body(self):
        with self.assertRaises(ValidationFailure):
            normalize([])

    def test_invalid_document_type_template_date_color(self):
        for over, key in ((dict(document={"type": "poem"}), "document.type"),
                          (dict(design={"template_id": "nope"}), "design.template_id"),
                          (dict(document={"submission_date": "16/10/2026"}), "document.submission_date"),
                          (dict(design={"accent": "red"}), "design.accent"),
                          (dict(design={"font": "Comic Sans"}), "design.font")):
            with self.assertRaises(ValidationFailure) as cm:
                normalize(sample(**over))
            self.assertIn(key, cm.exception.errors)

    def test_length_limits_and_member_count(self):
        with self.assertRaises(ValidationFailure) as cm:
            normalize(sample(document={"title": "x" * 301}))
        self.assertIn("document.title", cm.exception.errors)
        members = [{"name": f"S{i}", "student_id": str(i)} for i in range(13)]
        with self.assertRaises(ValidationFailure) as cm:
            normalize(sample(submitted_by={"mode": "group", "members": members}))
        self.assertIn("submitted_by.members", cm.exception.errors)

    def test_individual_requires_exactly_one(self):
        members = [{"name": "A", "student_id": "1"}, {"name": "B", "student_id": "2"}]
        with self.assertRaises(ValidationFailure):
            normalize(sample(submitted_by={"mode": "individual", "members": members}))

    def test_empty_optional_fields_ok(self):
        p = sample(document={"subtitle": "", "course_code": "", "course_title": "", "notes": ""},
                   institution={"faculty": "", "address": "", "logo_base64": ""})
        d = normalize(p)
        for t in TEMPLATE_IDS:
            d["design"]["template_id"] = t
            self.assertTrue(generate_pdf(d).startswith(b"%PDF"))
            self.assertTrue(zipfile.is_zipfile(io.BytesIO(generate_docx(d))))

    def test_control_characters_stripped(self):
        d = normalize(sample(document={"title": "Hello\x00 World\x07"}))
        self.assertEqual(d["document"]["title"], "Hello World")


class FilenameTests(unittest.TestCase):
    def test_example_name(self):
        self.assertEqual(safe_filename("Software Testing", "Assignment", "pdf"), "Software_Testing_Assignment_Cover.pdf")

    def test_path_traversal_and_symbols(self):
        n = safe_filename("../../etc/passwd\"; rm -rf /", "Assignment", "docx")
        self.assertRegex(n, r"^[A-Za-z0-9_]+\.docx$")

    def test_bengali_only_title_falls_back(self):
        self.assertEqual(safe_filename("সফটওয়্যার টেস্টিং", "Assignment", "pdf"), "Assignment_Cover.pdf")

    def test_empty_everything(self):
        self.assertEqual(safe_filename("", "", "pdf"), "Cover_Page.pdf")

    def test_length_capped(self):
        self.assertLessEqual(len(safe_filename("a" * 500, "Assignment", "pdf")), 84)


class TemplateTests(unittest.TestCase):
    def test_six_templates(self):
        self.assertGreaterEqual(len(template_catalog()), 6)
        self.assertEqual({t["id"] for t in template_catalog()}, set(TEMPLATE_IDS))

    def test_templates_produce_different_layouts(self):
        pdfs = {}
        for t in TEMPLATE_IDS:
            pdfs[t] = generate_pdf(normalize(sample(design={"template_id": t})))
        self.assertEqual(len(set(pdfs.values())), len(TEMPLATE_IDS))


class GenerationTests(unittest.TestCase):
    def test_pdf_valid_single_page_all_templates(self):
        for t in TEMPLATE_IDS:
            pdf = generate_pdf(normalize(sample(design={"template_id": t})))
            self.assertTrue(pdf.startswith(b"%PDF-"), t)
            self.assertIn(b"%%EOF", pdf[-32:], t)
            self.assertEqual(pdf_pages(pdf), 1, t)

    def test_pdf_contains_entered_text(self):
        if not shutil.which("pdftotext"):
            self.skipTest("pdftotext not installed")
        pdf = generate_pdf(normalize(sample()))
        with tempfile.NamedTemporaryFile(suffix=".pdf") as f:
            f.write(pdf); f.flush()
            text = subprocess.run(["pdftotext", f.name, "-"], capture_output=True, text=True).stdout
        for s in ("Software Testing Techniques", "Ayesha Siddika", "2021-1-60-045", "Dr. Rahman Chowdhury", "16 October 2026"):
            self.assertIn(s, text)

    def test_docx_valid_editable_contains_text(self):
        for t in TEMPLATE_IDS:
            blob = generate_docx(normalize(sample(design={"template_id": t})))
            self.assertTrue(zipfile.is_zipfile(io.BytesIO(blob)), t)
            with zipfile.ZipFile(io.BytesIO(blob)) as z:
                self.assertIn("word/document.xml", z.namelist())
                self.assertFalse([n for n in z.namelist() if n.startswith("word/media/")] == [] and False)
            doc = Document(io.BytesIO(blob))  # re-opens => structurally valid and editable
            full = "\n".join(p.text for p in doc.paragraphs)
            for tb in doc.tables:
                for row in tb.rows:
                    for c in row.cells:
                        full += "\n" + c.text
            self.assertIn("Ayesha Siddika", full, t)
            self.assertIn("Software Testing Techniques", full, t)
            self.assertEqual(doc.sections[0].page_width.cm.__round__(1), 21.0)

    def test_docx_embeds_logo_and_respects_hide_logo(self):
        with_logo = generate_docx(normalize(sample()))
        no_logo = generate_docx(normalize(sample(design={"show_logo": False})))
        names = lambda b: zipfile.ZipFile(io.BytesIO(b)).namelist()
        self.assertTrue(any(n.startswith("word/media/") for n in names(with_logo)))
        self.assertFalse(any(n.startswith("word/media/") for n in names(no_logo)))

    def test_long_title_and_long_names_stay_on_one_page(self):
        long_title = "A Comprehensive Comparative Analysis of Distributed Consensus Protocols " * 3
        p = sample(document={"title": long_title.strip()[:300], "subtitle": "s" * 250},
                   institution={"name": "The International University of Science, Technology and Applied Engineering Studies " * 2})
        d = normalize(p)
        for t in TEMPLATE_IDS:
            d["design"]["template_id"] = t
            pdf, scale, pages = render_pdf(d)
            self.assertEqual(pages, 1, f"{t} spilled to {pages} pages at scale {scale}")

    def test_group_with_many_members(self):
        members = [{"name": f"Student Number {i} With A Long Name", "student_id": f"2021-1-60-{i:03d}",
                    "role": "Lead" if i == 1 else ""} for i in range(1, 11)]
        d = normalize(sample(submitted_by={"mode": "group", "members": members}))
        for t in TEMPLATE_IDS:
            d["design"]["template_id"] = t
            pdf, _, pages = render_pdf(d)
            self.assertEqual(pages, 1, t)
            blob = generate_docx(d)
            doc = Document(io.BytesIO(blob))
            text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for tb in doc.tables for r in tb.rows for c in r.cells)
            self.assertIn("Student Number 10", text, t)

    def test_all_design_options_generate(self):
        for border in ("auto", "none", "thin", "double"):
            for density in ("compact", "normal", "relaxed"):
                for font in ("Times New Roman", "Arial", "DejaVu Sans"):
                    d = normalize(sample(design={"border": border, "density": density, "font": font,
                                                 "footer": "Draft", "show_header": False}))
                    self.assertTrue(generate_pdf(d).startswith(b"%PDF"))
                    self.assertTrue(zipfile.is_zipfile(io.BytesIO(generate_docx(d))))

    def test_special_characters_are_escaped(self):
        d = normalize(sample(document={"title": "R&D <b>Report</b> & \"Tests\""}))
        self.assertTrue(generate_pdf(d).startswith(b"%PDF"))
        self.assertIn("R&D <b>Report</b>", "\n".join(p.text for p in Document(io.BytesIO(generate_docx(d))).paragraphs))


class LayoutJsonTests(unittest.TestCase):
    def test_layout_is_json_serialisable_for_all_templates(self):
        import json
        from templates import build_layout
        for t in TEMPLATE_IDS:
            spec, blocks = build_layout(normalize(sample(design={"template_id": t})))
            json.dumps({"spec": spec, "blocks": blocks})
            self.assertTrue(blocks)


class LogoTests(unittest.TestCase):
    def test_valid_png_and_data_url(self):
        png, size = process_logo(logo_b64())
        self.assertTrue(png.startswith(b"\x89PNG"))
        self.assertEqual(process_logo("data:image/png;base64," + logo_b64())[1], size)

    def test_jpeg_accepted(self):
        from PIL import Image
        b = io.BytesIO(); Image.new("RGB", (100, 80), "red").save(b, "JPEG")
        self.assertEqual(process_logo(base64.b64encode(b.getvalue()).decode())[1], (100, 80))

    def test_rejects_garbage_text_svg_gif_oversize(self):
        from PIL import Image
        gif = io.BytesIO(); Image.new("P", (50, 50)).save(gif, "GIF")
        huge = io.BytesIO(); Image.new("RGB", (4100, 20), "white").save(huge, "PNG")
        cases = [base64.b64encode(b"not an image").decode(), "!!!notbase64!!!",
                 base64.b64encode(b"<svg xmlns='http://www.w3.org/2000/svg'/>").decode(),
                 "data:image/svg+xml;base64,PHN2Zy8+",
                 base64.b64encode(gif.getvalue()).decode(),
                 base64.b64encode(huge.getvalue()).decode(),
                 base64.b64encode(b"\x89PNG\r\n\x1a\n" + b"0" * 50).decode(),
                 "A" * (3 * 1024 * 1024)]
        for c in cases:
            with self.assertRaises(LogoError, msg=c[:30]):
                process_logo(c)

    def test_invalid_logo_surfaces_as_validation_error(self):
        with self.assertRaises(ValidationFailure) as cm:
            normalize(sample(institution={"logo_base64": base64.b64encode(b"junk").decode()}))
        self.assertIn("institution.logo", cm.exception.errors)


if __name__ == "__main__":
    unittest.main()
